#!/usr/bin/env python3
"""Apply the 《岩土工程学报》(CJGE) layout to an EXISTING .docx in place and audit it.

Covers everything except MathType equations: page setup (A4, CJGE margins, single-column front matter
and two-column body), front matter (题名, 作者, 单位, 摘要, 关键词, 中图分类号, 作者简介, English block,
first-page notes), headings, body, captions, 三线表, statements, references and the journal-style
running headers (page number in the header, no footer).

MathType content is out of scope and left untouched: runs holding OLE objects, MathType display
paragraphs (MTDisplayEquation / tab + Equation.DSMT object) and equation tables (tables whose cells hold
an Equation.DSMT object) are skipped. Their formatting belongs to the mathtype-for-word MCP.
Paragraphs that carry inline objects get "at least" instead of "exact" line spacing so inline
equations and pictures are never clipped.

The document tree is edited, never rebuilt: text, fields, bookmarks and objects keep their positions.

Usage:
  python apply_format.py apply INPUT.docx OUTPUT.docx [--profile P] [--header TEXT] [--year Y]
                         [--volume V] [--issue N] [--month M] [--no-journal-header] [--no-columns]
                         [--overwrite]
  python apply_format.py audit DOC.docx [--profile P]
"""
from __future__ import annotations

import argparse
import copy
import datetime
import json
import re
import sys
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PROFILE_DIR = Path(__file__).resolve().parent / "format_profiles"
DEFAULT_PROFILE = PROFILE_DIR / "cjge.json"

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
OBJECT_TAGS = (qn("w:object"), qn("w:drawing"), qn("w:pict"))
MATH_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
ALIGN = {"left": "left", "center": "center", "right": "right", "justify": "both", "both": "both"}
END_PUNCT = "。；;，,：:.!！?？"

HEADING_STYLE_RE = re.compile(r"^(heading|标题)\s*([1-9])$", re.I)
REFERENCE_HEADING_RE = re.compile(r"^\s*(参考文献|References|Bibliography)\s*[：:]?\s*$", re.I)
FIG_CN_RE = re.compile(r"^\s*图\s*[0-9]")
FIG_EN_RE = re.compile(r"^\s*(Fig\.?|Figure)\s*[0-9]", re.I)
TAB_CN_RE = re.compile(r"^\s*表\s*[0-9]")
TAB_EN_RE = re.compile(r"^\s*Table\s*[0-9]", re.I)
SUBCAP_RE = re.compile(r"^\s*[（(][a-z][)）]")
H1_TEXT_RE = re.compile(r"^\s*\d{1,2}(\s+|　)\S")
H2_TEXT_RE = re.compile(r"^\s*\d{1,2}\.\d{1,2}(\s+|　)\S")
H3_TEXT_RE = re.compile(r"^\s*[（(]\d{1,2}[）)]\s*\S")
FRONT = {
    "doi": re.compile(r"^\s*DOI\s*[：:]", re.I),
    "abstract_cn": re.compile(r"^\s*摘\s*要\s*[：:]"),
    "keywords_cn": re.compile(r"^\s*关\s*键\s*词\s*[：:]"),
    "clc": re.compile(r"^\s*(中图分类号|文献标识码|文章编号)\s*[：:]"),
    "bio": re.compile(r"^\s*作者简介\s*[：:]"),
    "abstract_en": re.compile(r"^\s*Abstract\s*[：:]", re.I),
    "keywords_en": re.compile(r"^\s*Key\s*words\s*[：:]", re.I),
    "footnote": re.compile(r"^\s*(基金项目\s*[：:]|\*?\s*通信作者|收稿日期)"),
}
STATEMENT_RE = re.compile(r"^\s*(利益冲突声明|作者贡献)")
SKIP_ROLES = ("empty", "mathtype", "figure")

PPR_ORDER = ["w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
             "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
             "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN",
             "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
             "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
             "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange"]
RPR_ORDER = ["w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike",
             "w:dstrike", "w:outline", "w:shadow", "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid",
             "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position", "w:sz", "w:szCs",
             "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd", "w:fitText", "w:vertAlign", "w:rtl", "w:cs",
             "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish", "w:oMath"]
SECTPR_ORDER = ["w:headerReference", "w:footerReference", "w:footnotePr", "w:endnotePr", "w:type", "w:pgSz",
                "w:pgMar", "w:paperSrc", "w:pgBorders", "w:lnNumType", "w:pgNumType", "w:cols", "w:formProt",
                "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid",
                "w:printerSettings", "w:sectPrChange"]


# --------------------------------------------------------------------------- profile

def load_profile(path: str = "") -> Dict[str, Any]:
    p = Path(path) if path else DEFAULT_PROFILE
    if not p.suffix:
        p = PROFILE_DIR / f"{p}.json"
    data = json.loads(p.read_text(encoding="utf-8-sig"))
    data["__path__"] = str(p.resolve())
    return data


def spec_for(profile: Dict[str, Any], role: str) -> Dict[str, Any]:
    roles = profile["roles"]
    if role in roles:
        return dict(roles[role])
    return dict(roles["body"])  # e.g. heading4+: 四级标题 a）… runs in with the body text


# --------------------------------------------------------------------------- classification

def _text(el) -> str:
    return "".join(t.text or "" for t in el.iter(qn("w:t")))


def _has_mathtype(el) -> bool:
    for ole in el.iter("{urn:schemas-microsoft-com:office:office}OLEObject"):
        if str(ole.get("ProgID", "")).startswith("Equation.DSMT"):
            return True
    return False


def _has_object(el) -> bool:
    return any(n.tag in OBJECT_TAGS or n.tag.startswith(MATH_NS + "oMath") for n in el.iter())


def _style_chain(paragraph) -> List[str]:
    names, style = [], paragraph.style
    while style is not None:
        names.append(style.name or "")
        style = style.base_style
    return names


def _short_heading(text: str, limit: int) -> bool:
    return 0 < len(text) <= limit and text[-1] not in END_PUNCT


def classify_text(paragraph, in_table: bool, in_refs: bool) -> str:
    """Role of one paragraph from its style and text alone (front matter is resolved afterwards)."""
    if in_table:
        return "table_body"
    styles = _style_chain(paragraph)
    text = _text(paragraph._p).strip()
    if "MTDisplayEquation" in styles or (_has_mathtype(paragraph._p) and not re.sub(r"[\s()（）0-9.\-，。,]", "", text)):
        return "mathtype"  # owned by the MathType workflow
    if not text:
        return "figure" if _has_object(paragraph._p) else "empty"
    if REFERENCE_HEADING_RE.match(text):
        return "reference_heading"
    for name in styles:
        if name.lower() in ("title", "标题"):
            return "title"
        m = HEADING_STYLE_RE.match(name)
        if m:
            return f"heading{m.group(2)}"
    if FIG_CN_RE.match(text):
        return "caption_fig_cn"
    if FIG_EN_RE.match(text):
        return "caption_fig_en"
    if TAB_CN_RE.match(text):
        return "caption_tab_cn"
    if TAB_EN_RE.match(text):
        return "caption_tab_en"
    if SUBCAP_RE.match(text) and len(text) <= 40:
        return "subcaption"
    if STATEMENT_RE.match(text):
        return "statement"
    if in_refs:
        return "reference"
    if H2_TEXT_RE.match(text) and _short_heading(text, 40):
        return "heading2"
    if H1_TEXT_RE.match(text) and _short_heading(text, 40):
        return "heading1"
    if H3_TEXT_RE.match(text) and _short_heading(text, 30):
        return "heading3"
    return "body"


def iter_body_paragraphs(document) -> List[Tuple[Any, str, bool]]:
    """(paragraph, role, top_level) in document order; equation tables are skipped entirely."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    out: List[Tuple[Any, str, bool]] = []
    state = {"refs": False}

    def walk(parent_el, parent_obj, in_table):
        for child in parent_el.iterchildren():
            if child.tag == qn("w:p"):
                paragraph = Paragraph(child, parent_obj)
                role = classify_text(paragraph, in_table, state["refs"])
                if not in_table:
                    if role == "reference_heading":
                        state["refs"] = True
                    elif role in ("title", "heading1", "heading2"):
                        state["refs"] = False
                out.append((paragraph, role, not in_table))
            elif child.tag == qn("w:tbl"):
                if _has_mathtype(child):
                    continue  # equation table: MathType workflow owns it
                table = Table(child, parent_obj)
                seen = set()
                for row in table.rows:
                    for cell in row.cells:
                        if id(cell._tc) in seen:
                            continue
                        seen.add(id(cell._tc))
                        walk(cell._tc, cell, True)
            elif child.tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    walk(content, parent_obj, in_table)

    walk(document.element.body, document, False)
    return out


def resolve_front_matter(items: List[Tuple[Any, str, bool]]) -> Dict[str, Any]:
    """Relabel the front matter (everything before the first level-1 heading) in place."""
    first_h1 = next((i for i, (_, r, top) in enumerate(items) if top and r == "heading1"), len(items))
    front = [i for i in range(first_h1) if items[i][2]]
    marks: Dict[str, int] = {}
    info: Dict[str, Any] = {"first_heading": first_h1, "front_end": None, "title": None, "authors": None}

    def setrole(i, role):
        p, old, top = items[i]
        if old not in SKIP_ROLES:
            items[i] = (p, role, top)

    for i in front:
        text = _text(items[i][0]._p)
        for role, rx in FRONT.items():
            if rx.match(text):
                marks.setdefault(role, i)
                setrole(i, role)
                break
    cn_start = marks.get("abstract_cn")
    if cn_start is not None:
        lead = [i for i in front if i < cn_start and items[i][1] not in SKIP_ROLES + ("doi", "footnote")]
        for k, i in enumerate(lead):
            setrole(i, "title" if k == 0 else "authors" if k == 1 else "affiliation")
        if lead:
            info["title"] = _text(items[lead[0]][0]._p).strip()
        if len(lead) > 1:
            info["authors"] = _text(items[lead[1]][0]._p).strip()
        end_cn = marks.get("keywords_cn", cn_start)
        for i in front:
            if cn_start < i < end_cn and items[i][1] == "body":
                setrole(i, "abstract_cn")
    else:
        titled = [i for i in front if items[i][1] == "title"]
        if titled:
            info["title"] = _text(items[titled[0]][0]._p).strip()
    en_start = marks.get("abstract_en")
    if en_start is not None:
        after_cn = max([marks[k] for k in ("keywords_cn", "clc", "bio", "abstract_cn") if k in marks] or [-1])
        lead = [i for i in front if after_cn < i < en_start and items[i][1] not in SKIP_ROLES + (
            "footnote", "clc", "bio")]
        for k, i in enumerate(lead):
            setrole(i, "title_en" if k == 0 else "authors_en" if k == 1 else "affiliation_en")
        end_en = marks.get("keywords_en", en_start)
        for i in front:
            if en_start < i < end_en and items[i][1] == "body":
                setrole(i, "abstract_en")
    block = [v for k, v in marks.items() if k not in ("footnote", "doi")]
    if block:
        info["front_end"] = max(block)
    return info


def classify_document(document) -> Tuple[List[Tuple[Any, str, bool]], Dict[str, Any]]:
    items = iter_body_paragraphs(document)
    return items, resolve_front_matter(items)


# --------------------------------------------------------------------------- xml helpers

def _child(parent, tag):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def _reorder(el, order):
    rank = {qn(t): i for i, t in enumerate(order)}
    kids = sorted(list(el), key=lambda c: rank.get(c.tag, len(order)))
    for c in kids:
        el.remove(c)
        el.append(c)


def set_rpr(rPr, spec) -> None:
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), spec["latin"])
    rf.set(qn("w:eastAsia"), spec["east"])
    if spec.get("size"):
        half = str(int(round(float(spec["size"]) * 2)))
        _child(rPr, "w:sz").set(qn("w:val"), half)
        _child(rPr, "w:szCs").set(qn("w:val"), half)
    _child(rPr, "w:lang").set(qn("w:eastAsia"), "zh-CN")
    if spec.get("bold") is not None:
        for tag in ("w:b", "w:bCs"):
            _child(rPr, tag).set(qn("w:val"), "1" if spec["bold"] else "0")
    if spec.get("color"):
        color = _child(rPr, "w:color")
        for attr in list(color.attrib):
            del color.attrib[attr]
        color.set(qn("w:val"), spec["color"])
    _reorder(rPr, RPR_ORDER)


def line_setting(spec, has_object: bool = False) -> Tuple[str, str]:
    """(lineRule, line) for a spec; exact spacing becomes atLeast when the paragraph holds objects."""
    rule = spec.get("line_rule", "auto")
    if rule in ("exact", "at_least"):
        word_rule = "atLeast" if rule == "at_least" or has_object else "exact"
        return word_rule, str(int(round(float(spec["line_pt"]) * 20)))
    return "auto", str(int(round(240 * float(spec.get("line_multiple") or 1.0))))


def set_ppr(pPr, spec, has_object: bool = False) -> None:
    if spec.get("align"):
        _child(pPr, "w:jc").set(qn("w:val"), ALIGN[spec["align"]])
    ind = _child(pPr, "w:ind")
    for attr in list(ind.attrib):
        del ind.attrib[attr]
    ind.set(qn("w:left"), str(int(round(float(spec.get("left_pt") or 0) * 20))))
    ind.set(qn("w:right"), str(int(round(float(spec.get("right_pt") or 0) * 20))))
    if spec.get("hanging_pt"):
        ind.set(qn("w:hanging"), str(int(round(float(spec["hanging_pt"]) * 20))))
    elif spec.get("first_chars"):
        ind.set(qn("w:firstLineChars"), str(int(round(float(spec["first_chars"]) * 100))))
        ind.set(qn("w:firstLine"), str(int(round(float(spec["first_chars"]) * float(spec["size"]) * 20))))
    else:
        ind.set(qn("w:firstLineChars"), "0")
        ind.set(qn("w:firstLine"), str(int(round(float(spec.get("first_pt") or 0) * 20))))
    sp = _child(pPr, "w:spacing")
    for attr in list(sp.attrib):
        del sp.attrib[attr]
    sp.set(qn("w:before"), str(int(round(float(spec.get("before") or 0) * 20))))
    sp.set(qn("w:beforeLines"), "0")
    sp.set(qn("w:after"), str(int(round(float(spec.get("after") or 0) * 20))))
    sp.set(qn("w:afterLines"), "0")
    rule, line = line_setting(spec, has_object)
    sp.set(qn("w:line"), line)
    sp.set(qn("w:lineRule"), rule)
    _reorder(pPr, PPR_ORDER)


def _run_is_object(r) -> bool:
    return any(n.tag in OBJECT_TAGS for n in r.iter())


def _paragraph_runs(p_el) -> List[Any]:
    """Runs of a paragraph in reading order (also inside hyperlinks, fields, insertions)."""
    return list(p_el.iter(qn("w:r")))


def label_ranges(text: str, spec) -> List[Tuple[int, int, Dict[str, Any]]]:
    ranges = []
    for label in spec.get("labels", []):
        for m in re.finditer(label["pattern"], text):
            if m.end() > m.start():
                override = {k: v for k, v in label.items() if k != "pattern"}
                start = m.start() + (len(m.group(0)) - len(m.group(0).lstrip()))
                ranges.append((start, m.end(), override))
    return ranges


def _split_run(r, cut: int) -> None:
    """Split a run that holds only w:t text at character offset `cut`."""
    kids = [k for k in r if k.tag != qn("w:rPr")]
    if not kids or any(k.tag != qn("w:t") for k in kids):
        return
    text = "".join(k.text or "" for k in kids)
    if not 0 < cut < len(text):
        return
    for k in kids[1:]:
        r.remove(k)
    kids[0].text = text[:cut]
    kids[0].set(XML_SPACE, "preserve")
    right = copy.deepcopy(r)
    right.find(qn("w:t")).text = text[cut:]
    r.addnext(right)


def split_at_labels(p_el, ranges) -> None:
    for cut in sorted({c for a, b, _ in ranges for c in (a, b)}):
        pos = 0
        for r in _paragraph_runs(p_el):
            t = _text(r)
            if pos < cut < pos + len(t):
                _split_run(r, cut - pos)
                break
            pos += len(t)


def run_specs(p_el, spec) -> List[Tuple[Any, Dict[str, Any]]]:
    """(run, expected spec) for every run, with label overrides applied by character offset."""
    ranges = label_ranges(_text(p_el), spec)
    base = {k: v for k, v in spec.items() if k != "labels"}
    out, pos = [], 0
    for r in _paragraph_runs(p_el):
        merged = dict(base)
        for a, b, override in ranges:
            if a <= pos < b:
                merged.update(override)
        out.append((r, merged))
        pos += len(_text(r))
    return out


def format_paragraph(p_el, spec) -> int:
    pPr = p_el.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_el.insert(0, pPr)
    set_ppr(pPr, spec, has_object=_has_object(p_el))
    ranges = label_ranges(_text(p_el), spec)
    if ranges:
        split_at_labels(p_el, ranges)
    touched = 0
    for r, run_spec in run_specs(p_el, spec):
        if _run_is_object(r):
            continue  # never touch runs that hold objects (MathType OLE, pictures)
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            rPr = OxmlElement("w:rPr")
            r.insert(0, rPr)
        set_rpr(rPr, run_spec)
        touched += 1
    return touched


# --------------------------------------------------------------------------- tables

TBLPR_ORDER = ["w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
               "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
               "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription"]
TCPR_ORDER = ["w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd", "w:noWrap",
              "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"]


def _border(parent, side: str, size_eighths: int, val: str = "single") -> None:
    el = _child(parent, f"w:{side}")
    for attr in list(el.attrib):
        del el.attrib[attr]
    if size_eighths:
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(size_eighths))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
    else:
        el.set(qn("w:val"), "nil")


def format_three_line_table(tbl, cfg: Dict[str, Any], body_spec: Dict[str, Any]) -> None:
    """三线表: top and bottom rules, one thinner rule under the header row, no other borders."""
    thick = int(round(float(cfg.get("top_bottom_pt", 0.75)) * 8))
    thin = int(round(float(cfg.get("header_rule_pt", 0.5)) * 8))
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    style = tblPr.find(qn("w:tblStyle"))
    if style is not None:
        tblPr.remove(style)  # table styles (e.g. Table Grid) would draw their own borders and fills
    if cfg.get("align"):
        _child(tblPr, "w:jc").set(qn("w:val"), cfg["align"])
    borders = _child(tblPr, "w:tblBorders")
    for side, size in (("top", thick), ("left", 0), ("bottom", thick), ("right", 0),
                       ("insideH", 0), ("insideV", 0)):
        _border(borders, side, size)
    _reorder(borders, ["w:top", "w:left", "w:start", "w:bottom", "w:right", "w:end", "w:insideH", "w:insideV"])
    mar = _child(tblPr, "w:tblCellMar")
    for side, key in (("top", "cell_margin_tb_twips"), ("left", "cell_margin_lr_twips"),
                      ("bottom", "cell_margin_tb_twips"), ("right", "cell_margin_lr_twips")):
        el = _child(mar, f"w:{side}")
        el.set(qn("w:w"), str(int(cfg.get(key, 0))))
        el.set(qn("w:type"), "dxa")
    _reorder(mar, ["w:top", "w:left", "w:start", "w:bottom", "w:right", "w:end"])
    _reorder(tblPr, TBLPR_ORDER)
    rows = tbl.findall(qn("w:tr"))
    for r_index, tr in enumerate(rows):
        for tc in tr.findall(qn("w:tc")):
            tcPr = tc.find(qn("w:tcPr"))
            if tcPr is None:
                tcPr = OxmlElement("w:tcPr")
                tc.insert(0, tcPr)
            old = tcPr.find(qn("w:tcBorders"))
            if old is not None:
                tcPr.remove(old)
            shd = tcPr.find(qn("w:shd"))
            if shd is not None:
                tcPr.remove(shd)  # no fill colour
            if r_index == 0 and len(rows) > 1:
                _border(_child(tcPr, "w:tcBorders"), "bottom", thin)
            _reorder(tcPr, TCPR_ORDER)
            if r_index == 0:
                for p in tc.iter(qn("w:p")):
                    for r in _paragraph_runs(p):
                        if _run_is_object(r):
                            continue
                        rPr = r.find(qn("w:rPr"))
                        if rPr is None:
                            rPr = OxmlElement("w:rPr")
                            r.insert(0, rPr)
                        set_rpr(rPr, {**body_spec, "bold": bool(cfg.get("header_bold"))})


def audit_three_line_table(tbl, cfg: Dict[str, Any]) -> List[str]:
    problems = []
    thick = str(int(round(float(cfg.get("top_bottom_pt", 0.75)) * 8)))
    thin = str(int(round(float(cfg.get("header_rule_pt", 0.5)) * 8)))
    borders = tbl.find(f"{qn('w:tblPr')}/{qn('w:tblBorders')}")
    if borders is None:
        return ["missing table borders"]
    for side, want in (("top", thick), ("bottom", thick), ("left", None), ("right", None),
                       ("insideH", None), ("insideV", None)):
        el = borders.find(qn(f"w:{side}"))
        val = None if el is None else el.get(qn("w:val"))
        if want is None and val not in ("nil", "none"):
            problems.append(f"{side} border should be none (is {val})")
        if want is not None and (val != "single" or el.get(qn("w:sz")) != want):
            problems.append(f"{side} border should be single {want}/8 pt")
    if tbl.find(f"{qn('w:tblPr')}/{qn('w:tblStyle')}") is not None:
        problems.append("table style still applied")
    rows = tbl.findall(qn("w:tr"))
    if len(rows) > 1:
        for tc in rows[0].findall(qn("w:tc")):
            b = tc.find(f"{qn('w:tcPr')}/{qn('w:tcBorders')}/{qn('w:bottom')}")
            if b is None or b.get(qn("w:sz")) != thin:
                problems.append("header row lacks the thin bottom rule")
                break
    return problems


# --------------------------------------------------------------------------- page / sections

def _twips(cm: float) -> str:
    return str(int(round(cm / 2.54 * 1440)))


def set_page(sectPr, page: Dict[str, Any], columns: int) -> None:
    pgSz = _child(sectPr, "w:pgSz")
    pgSz.set(qn("w:w"), _twips(page["width_cm"]))
    pgSz.set(qn("w:h"), _twips(page["height_cm"]))
    pgSz.attrib.pop(qn("w:orient"), None)
    pgMar = _child(sectPr, "w:pgMar")
    for attr, key in (("w:top", "top_cm"), ("w:bottom", "bottom_cm"), ("w:left", "left_cm"),
                      ("w:right", "right_cm"), ("w:header", "header_cm"), ("w:footer", "footer_cm")):
        pgMar.set(qn(attr), _twips(page[key]))
    pgMar.set(qn("w:gutter"), "0")
    cols = _child(sectPr, "w:cols")
    for attr in list(cols.attrib):
        del cols.attrib[attr]
    for c in list(cols):
        cols.remove(c)
    cols.set(qn("w:num"), str(columns))
    cols.set(qn("w:space"), _twips(page.get("column_space_cm", 0.74)))
    _reorder(sectPr, SECTPR_ORDER)


def apply_sections(document, page: Dict[str, Any], front_end_p, use_columns: bool) -> Dict[str, Any]:
    """Front matter single column, body in `body_columns` columns (continuous section break)."""
    body = document.element.body
    final = body.find(qn("w:sectPr"))
    inner = [p for p in body.iter(qn("w:p")) if p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None]
    n_body = int(page.get("body_columns", 1)) if use_columns else 1
    result = {"sections": 1 + len(inner), "front_matter_break": False}
    if use_columns and front_end_p is not None and not inner and final is not None and n_body > 1:
        pPr = front_end_p.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            front_end_p.insert(0, pPr)
        first = copy.deepcopy(final)
        for ref in final.findall(qn("w:headerReference")) + final.findall(qn("w:footerReference")):
            final.remove(ref)  # headers now belong to the first section; the body section inherits them
        titlepg = final.find(qn("w:titlePg"))
        if titlepg is not None:
            final.remove(titlepg)
        pPr.append(first)
        _reorder(pPr, PPR_ORDER)
        set_page(first, page, 1)
        _child(final, "w:type").set(qn("w:val"), "continuous")
        _reorder(final, SECTPR_ORDER)
        set_page(final, page, n_body)
        result.update(sections=2, front_matter_break=True)
        return result
    all_sect = [p.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") for p in inner] + ([final] if final is not None else [])
    for i, s in enumerate(all_sect):
        single_front = use_columns and i == 0 and len(all_sect) > 1 and front_end_p is not None
        set_page(s, page, 1 if single_front else n_body)
    return result


# --------------------------------------------------------------------------- headers

def _field_runs(instr: str, spec) -> List[Any]:
    runs = []
    for kind, value in (("begin", None), ("instr", f" {instr} "), ("separate", None), ("text", "1"), ("end", None)):
        r = OxmlElement("w:r")
        rPr = OxmlElement("w:rPr")
        r.append(rPr)
        set_rpr(rPr, spec)
        if kind in ("begin", "separate", "end"):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r.append(fc)
        elif kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(XML_SPACE, "preserve")
            it.text = value
            r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = value
            r.append(t)
        runs.append(r)
    return runs


def _text_run(text: str, spec) -> Any:
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    r.append(rPr)
    set_rpr(rPr, spec)
    t = OxmlElement("w:t")
    t.set(XML_SPACE, "preserve")
    t.text = text
    r.append(t)
    return r


def _tab_run(spec) -> Any:
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    r.append(rPr)
    set_rpr(rPr, spec)
    r.append(OxmlElement("w:tab"))
    return r


def _header_paragraph(parts: List[Any], spec, width: int, rule: Optional[Tuple[str, int]]) -> Any:
    """Paragraph with left | centre | right parts separated by centre and right tab stops.

    A part is a string, ("PAGE",) for a page-number field, or (overrides, text) for styled text."""
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    p.append(pPr)
    tabs = OxmlElement("w:tabs")
    for val, pos in (("center", width // 2), ("right", width)):
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), val)
        tab.set(qn("w:pos"), str(pos))
        tabs.append(tab)
    pPr.append(tabs)
    if rule:
        pBdr = OxmlElement("w:pBdr")
        _border(pBdr, "bottom", rule[1], rule[0])
        pBdr.find(qn("w:bottom")).set(qn("w:space"), "1")
        pPr.append(pBdr)
    set_ppr(pPr, {**spec, "align": "left", "first_pt": 0})
    for k, part in enumerate(parts):
        if k:
            p.append(_tab_run(spec))
        if isinstance(part, tuple) and part and part[0] == "PAGE":
            p.extend(_field_runs("PAGE", spec))
        elif isinstance(part, tuple):
            p.append(_text_run(part[1], {**spec, **part[0]}))
        elif part:
            p.append(_text_run(part, spec))
    return p


def _fill(part_obj, paragraphs: List[Any]) -> None:
    el = part_obj._element
    for child in list(el):
        if child.tag in (qn("w:p"), qn("w:tbl"), qn("w:sdt")):
            el.remove(child)
    for p in paragraphs or [OxmlElement("w:p")]:
        el.append(p)


def running_header_text(info: Dict[str, Any]) -> str:
    """Odd-page header: 第一作者，等. 中文题名 (or 作者. 题名 for a single author)."""
    title = (info.get("title") or "").strip()
    authors = (info.get("authors") or "").strip()
    names = [re.sub(r"[\d*＊\s]", "", n) for n in re.split(r"[，、]|,\s*(?=\D)", authors)]
    names = [n for n in names if n]
    if names and title:
        return f"{names[0]}，等. {title}" if len(names) > 1 else f"{names[0]}. {title}"
    return title


def apply_journal_header(document, prof, info, header_text: Optional[str], year: int, volume: str, issue: str,
                         month: str) -> Dict[str, Any]:
    cfg = prof.get("journal_header", {})
    page = prof["page"]
    width = int(_twips(page["width_cm"] - page["left_cm"] - page["right_cm"]))
    spec = spec_for(prof, "header")
    rule = ("single", int(round(float(cfg.get("rule_pt", 0.5)) * 8)))
    odd_center = header_text if header_text is not None else running_header_text(info)
    issue_txt = f"第{issue}期" if issue else ""
    document.settings.odd_and_even_pages_header_footer = True
    sections = document.sections
    first_section = sections[0]
    first_section.different_first_page_header_footer = True
    first_section.header.is_linked_to_previous = False
    _fill(first_section.header, [_header_paragraph([issue_txt, odd_center, ("PAGE",)], spec, width, rule)])
    first_section.even_page_header.is_linked_to_previous = False
    _fill(first_section.even_page_header,
          [_header_paragraph([("PAGE",), cfg.get("journal_cn", ""), f"{year}年"], spec, width, rule)])
    first_page = []
    if volume and issue:
        months = ["Jan.", "Feb.", "Mar.", "Apr.", "May", "Jun.", "Jul.", "Aug.", "Sep.", "Oct.", "Nov.", "Dec."]
        m = int(month) if str(month).isdigit() else 0
        mon_en = months[m - 1] if 1 <= m <= 12 else ""
        big = spec_for(prof, "header_journal")
        double = (cfg.get("first_page_rule", "thickThinSmallGap"),
                  int(round(float(cfg.get("first_page_rule_pt", 2.25)) * 8)))
        first_page = [
            _header_paragraph([f"第{volume}卷 第{issue}期", ({"size": big["size"]}, cfg.get("journal_cn", "")),
                               f"Vol.{volume} No.{issue}"], spec, width, None),
            _header_paragraph([f"{year}年 {m}月" if m else f"{year}年", cfg.get("journal_en", ""),
                               f"{mon_en} {year}".strip()], spec, width, double),
        ]
    first_section.first_page_header.is_linked_to_previous = False
    _fill(first_section.first_page_header, first_page)
    for kind in ("footer", "even_page_footer", "first_page_footer"):
        footer = getattr(first_section, kind)
        footer.is_linked_to_previous = False
        _fill(footer, [])  # CJGE: no footer, the page number lives in the header
    for section in sections[1:]:
        # Later sections inherit the running headers. References are removed directly: python-docx's
        # unlink would also drop relationships that the first section still uses.
        sp = section._sectPr
        for ref in sp.findall(qn("w:headerReference")) + sp.findall(qn("w:footerReference")):
            sp.remove(ref)
        section.different_first_page_header_footer = False
    return {"odd_header": odd_center, "even_header": cfg.get("journal_cn", ""), "first_page_header": bool(first_page),
            "footer": "none"}


# --------------------------------------------------------------------------- footnotes

def format_footnotes(document, prof) -> int:
    """First-page notes (基金项目, 通信作者, 收稿日期) kept as Word footnotes get the footnote role."""
    from lxml import etree
    count = 0
    spec = spec_for(prof, "footnote")
    for rel in document.part.rels.values():
        if not rel.reltype.endswith("/footnotes") or rel.is_external:
            continue
        part = rel.target_part
        is_xml = hasattr(part, "element")
        root = part.element if is_xml else etree.fromstring(part.blob)
        for fn in root.findall(qn("w:footnote")):
            if fn.get(qn("w:type")) in ("separator", "continuationSeparator", "continuationNotice"):
                continue
            for p in fn.findall(qn("w:p")):
                format_paragraph(p, spec)
                count += 1
        if not is_xml:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    return count


# --------------------------------------------------------------------------- styles

def apply_styles(document, prof) -> None:
    """Defaults and styles, so text typed later inherits the CJGE format."""
    rpr_default = document.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
    if rpr_default is not None:
        set_rpr(rpr_default, spec_for(prof, "body"))
    for style in document.styles:
        name = style.name or ""
        m = HEADING_STYLE_RE.match(name)
        role = "body" if name == "Normal" else "title" if name.lower() in ("title", "标题") else (
            "caption_fig_cn" if name.lower() in ("caption", "题注") else
            f"heading{m.group(2)}" if m and int(m.group(2)) <= 3 else None)
        if not role:
            continue
        spec = spec_for(prof, role)
        el = style.element
        rpr = el.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            el.append(rpr)
        set_rpr(rpr, spec)
        ppr = el.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            rpr.addprevious(ppr)
        set_ppr(ppr, spec)


# --------------------------------------------------------------------------- apply

def apply(input_path: str, output_path: str, profile_path: str = "", header_text: Optional[str] = None,
          overwrite: bool = False, year: Optional[int] = None, volume: str = "", issue: str = "",
          month: str = "", journal_header: bool = True, columns: bool = True) -> Dict[str, Any]:
    src, dst = Path(input_path).resolve(), Path(output_path).resolve()
    if src == dst:
        raise ValueError("output must differ from input; the source is preserved")
    if dst.exists() and not overwrite:
        raise FileExistsError(f"output exists: {dst} (use --overwrite)")
    prof = load_profile(profile_path)
    document = docx.Document(str(src))
    apply_styles(document, prof)

    items, info = classify_document(document)
    counts: Dict[str, int] = {}
    runs = 0
    for paragraph, role, _top in items:
        counts[role] = counts.get(role, 0) + 1
        if role != "mathtype":
            runs += format_paragraph(paragraph._p, spec_for(prof, role))

    tables = 0
    cfg = prof.get("three_line_tables")
    if cfg:
        for tbl in document.element.body.iter(qn("w:tbl")):
            if not _has_mathtype(tbl):
                format_three_line_table(tbl, cfg, spec_for(prof, "table_body"))
                tables += 1

    front_end_p = items[info["front_end"]][0]._p if info.get("front_end") is not None else None
    sections = apply_sections(document, prof["page"], front_end_p, columns)
    footnotes = format_footnotes(document, prof)

    if journal_header:
        hf = apply_journal_header(document, prof, info, header_text, year or datetime.date.today().year,
                                  volume, issue, month)
    else:
        hf = {"header_paragraphs": 0, "footer_paragraphs": 0}
        page = prof["page"]
        width = int(_twips(page["width_cm"] - page["left_cm"] - page["right_cm"]))
        for section in document.sections:
            if header_text is not None:
                section.header.is_linked_to_previous = False
                _fill(section.header, [_header_paragraph(["", header_text], spec_for(prof, "header"), width, None)])
            for kind in ("header", "footer"):
                for p in getattr(section, kind).paragraphs:
                    if _text(p._p).strip() or p._p.find(f".//{qn('w:fldChar')}") is not None:
                        format_paragraph(p._p, spec_for(prof, kind))
                        hf[f"{kind}_paragraphs"] += 1

    tmp = dst.with_name(f".{dst.stem}.tmp{dst.suffix}")
    document.save(str(tmp))
    tmp.replace(dst)
    return {"ok": True, "input": str(src), "output": str(dst), "profile": prof["__path__"],
            "paragraph_roles": counts, "runs_formatted": runs, "three_line_tables": tables,
            "sections": sections, "footnote_paragraphs": footnotes, "header_footer": hf,
            "front_matter": {"title": info.get("title"), "authors": info.get("authors")}}


# --------------------------------------------------------------------------- audit

class EffectiveFormat:
    """Resolves paragraph/run properties the way Word does: direct formatting, then the character
    style chain, the paragraph style chain and finally the document defaults."""

    def __init__(self, document):
        root = document.styles.element
        self.styles = {s.get(qn("w:styleId")): s for s in root.findall(qn("w:style"))}
        self.default_para = next((sid for sid, s in self.styles.items() if s.get(qn("w:type")) == "paragraph"
                                  and s.get(qn("w:default")) in ("1", "true")), None)
        self.ppr_default = root.find(f"{qn('w:docDefaults')}/{qn('w:pPrDefault')}/{qn('w:pPr')}")
        self.rpr_default = root.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")

    def _style_chain(self, style_id, tag):
        out, seen = [], set()
        while style_id and style_id in self.styles and style_id not in seen:
            seen.add(style_id)
            style = self.styles[style_id]
            el = style.find(qn(tag))
            if el is not None:
                out.append(el)
            based = style.find(qn("w:basedOn"))
            style_id = based.get(qn("w:val")) if based is not None else None
        return out

    def _pstyle(self, p_el):
        ps = p_el.find(f"{qn('w:pPr')}/{qn('w:pStyle')}")
        return ps.get(qn("w:val")) if ps is not None else self.default_para

    def paragraph_chain(self, p_el):
        chain = [p_el.find(qn("w:pPr"))] + self._style_chain(self._pstyle(p_el), "w:pPr") + [self.ppr_default]
        return [c for c in chain if c is not None]

    def run_chain(self, r, p_el):
        rpr = r.find(qn("w:rPr"))
        rstyle = rpr.find(qn("w:rStyle")) if rpr is not None else None
        chain = [rpr] + (self._style_chain(rstyle.get(qn("w:val")), "w:rPr") if rstyle is not None else [])
        chain += self._style_chain(self._pstyle(p_el), "w:rPr") + [self.rpr_default]
        return [c for c in chain if c is not None]

    @staticmethod
    def element(chain, tag):
        for props in chain:
            el = props.find(qn(tag))
            if el is not None:
                return el
        return None

    @staticmethod
    def value(chain, tag, attr):
        for props in chain:
            el = props.find(qn(tag))
            if el is not None and el.get(qn(attr)) is not None:
                return el.get(qn(attr))
        return None


def audit(docx_path: str, profile_path: str = "") -> Dict[str, Any]:
    prof = load_profile(profile_path)
    document = docx.Document(str(Path(docx_path).resolve()))
    issues: List[Dict[str, Any]] = []
    checked = 0

    eff = EffectiveFormat(document)

    def check(p_el, spec, where):
        pchain = eff.paragraph_chain(p_el)
        rule, line = line_setting(spec, _has_object(p_el))
        got_rule = eff.value(pchain, "w:spacing", "w:lineRule") or "auto"
        got_line = eff.value(pchain, "w:spacing", "w:line")
        if got_rule != rule or got_line != line:
            issues.append({**where, "issue": "line_spacing", "expected": f"{rule}:{line}",
                           "actual": f"{got_rule}:{got_line}"})
        jc = eff.value(pchain, "w:jc", "w:val")
        if spec.get("align") and jc != ALIGN[spec["align"]]:
            issues.append({**where, "issue": "alignment", "expected": ALIGN[spec["align"]], "actual": jc})
        wants = []
        if spec.get("first_chars"):
            wants.append(("w:firstLine", str(int(round(float(spec["first_chars"]) * float(spec["size"]) * 20)))))
        for key, attr in (("left_pt", "w:left"), ("right_pt", "w:right"), ("hanging_pt", "w:hanging")):
            if spec.get(key):
                wants.append((attr, str(int(round(float(spec[key]) * 20)))))
        for attr, want in wants:
            got = eff.value(pchain, "w:ind", attr)
            if got != want:
                issues.append({**where, "issue": f"indent {attr[2:]}", "expected": want, "actual": got})
        for r, run_spec in run_specs(p_el, spec):
            if _run_is_object(r) or not _text(r).strip():
                continue
            rchain = eff.run_chain(r, p_el)
            ascii_font = eff.value(rchain, "w:rFonts", "w:ascii")
            east_font = eff.value(rchain, "w:rFonts", "w:eastAsia")
            if ascii_font != run_spec["latin"] or east_font != run_spec["east"]:
                issues.append({**where, "issue": "font", "expected": f"{run_spec['latin']}/{run_spec['east']}",
                               "actual": f"{ascii_font}/{east_font}", "run": _text(r)[:12]})
                break
            sz = eff.value(rchain, "w:sz", "w:val")
            if run_spec.get("size") and (sz is None or float(sz) / 2 != float(run_spec["size"])):
                issues.append({**where, "issue": "font_size", "expected": run_spec["size"],
                               "actual": None if sz is None else float(sz) / 2, "run": _text(r)[:12]})
                break
            if run_spec.get("bold") is not None:
                b = eff.element(rchain, "w:b")
                is_bold = b is not None and b.get(qn("w:val"), "1") not in ("0", "false")
                if is_bold != bool(run_spec["bold"]):
                    issues.append({**where, "issue": "bold", "expected": run_spec["bold"], "actual": is_bold,
                                   "run": _text(r)[:12]})
                    break

    items, _info = classify_document(document)
    for index, (paragraph, role, _top) in enumerate(items, 1):
        if role == "mathtype":
            continue
        checked += 1
        check(paragraph._p, spec_for(prof, role),
              {"paragraph": index, "role": role, "text": _text(paragraph._p).strip()[:30]})
    cfg = prof.get("three_line_tables")
    if cfg:
        for t_i, tbl in enumerate(document.element.body.iter(qn("w:tbl"))):
            if _has_mathtype(tbl):
                continue
            checked += 1
            for problem in audit_three_line_table(tbl, cfg):
                issues.append({"table": t_i, "role": "three_line_table", "issue": problem})
    page = prof["page"]
    sections = document.sections
    for s_i, section in enumerate(sections):
        sp = section._sectPr
        pgMar = sp.find(qn("w:pgMar"))
        for attr, key in (("w:left", "left_cm"), ("w:right", "right_cm"), ("w:top", "top_cm"),
                          ("w:bottom", "bottom_cm")):
            if pgMar is None or pgMar.get(qn(attr)) != _twips(page[key]):
                issues.append({"section": s_i, "role": "page", "issue": f"margin {attr[2:]}",
                               "expected": _twips(page[key]), "actual": None if pgMar is None else pgMar.get(qn(attr))})
        if s_i == len(sections) - 1:
            cols = sp.find(qn("w:cols"))
            num = "1" if cols is None else cols.get(qn("w:num"), "1")
            if num != str(page.get("body_columns", 1)):
                issues.append({"section": s_i, "role": "page", "issue": "body columns",
                               "expected": page.get("body_columns"), "actual": num})
        for kind in ("header", "even_page_header", "first_page_header", "footer"):
            part = getattr(section, kind)
            if part.is_linked_to_previous:
                continue
            hspec = spec_for(prof, "footer" if kind == "footer" else "header")
            for p in part.paragraphs:
                for r in _paragraph_runs(p._p):
                    rchain = eff.run_chain(r, p._p)
                    if _text(r).strip() and (eff.value(rchain, "w:rFonts", "w:ascii") != hspec["latin"]
                                             or eff.value(rchain, "w:rFonts", "w:eastAsia") != hspec["east"]):
                        issues.append({"section": s_i, "role": kind, "issue": "font", "text": _text(p._p)[:30]})
                        break
                if _text(p._p).strip():
                    checked += 1
    return {"ok": not issues, "document": str(Path(docx_path).resolve()), "profile": prof["__path__"],
            "paragraphs_checked": checked, "issue_count": len(issues), "issues": issues[:200]}


def main() -> int:
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("apply")
    a.add_argument("input")
    a.add_argument("output")
    a.add_argument("--profile", default="")
    a.add_argument("--header", default=None, help="odd-page running header; default: 第一作者，等. 题名")
    a.add_argument("--year", type=int, default=None)
    a.add_argument("--volume", default="")
    a.add_argument("--issue", default="")
    a.add_argument("--month", default="")
    a.add_argument("--no-journal-header", action="store_true", help="keep the existing header/footer layout")
    a.add_argument("--no-columns", action="store_true", help="do not switch the body to two columns")
    a.add_argument("--overwrite", action="store_true")
    b = sub.add_parser("audit")
    b.add_argument("docx")
    b.add_argument("--profile", default="")
    args = ap.parse_args()
    if args.cmd == "apply":
        result = apply(args.input, args.output, args.profile, args.header, args.overwrite, args.year,
                       args.volume, args.issue, args.month, not args.no_journal_header, not args.no_columns)
    else:
        result = audit(args.docx, args.profile)
    sys.stdout.reconfigure(encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=1))
    return 0 if result.get("ok") else 1


if __name__ == "__main__":
    raise SystemExit(main())
